from pptx import Presentation
from pptx.util import Pt

prs = Presentation()

# Slide 1: Title
slide_1 = prs.slides.add_slide(prs.slide_layouts[0])
slide_1.shapes.title.text = "IoT-Based Smart Parking Occupancy Monitoring & Predictive Demand System"
slide_1.placeholders[1].text = (
    "Project Proposal\n\n"
    "Group Members\n"
    "1____________________________ | ____________________________\n"
    "2____________________________ | ____________________________\n"
    "3____________________________ | ____________________________\n"
    "4____________________________ | ____________________________\n"
    "5____________________________ | ____________________________\n\n"
    "BSc Telecommunication & Information Engineering\n"
    "Dedan Kimathi University of Technology"
)
for paragraph in slide_1.placeholders[1].text_frame.paragraphs:
    paragraph.font.size = Pt(16)

# Slide 2: Problem Statement
slide_2 = prs.slides.add_slide(prs.slide_layouts[1])
slide_2.shapes.title.text = "Problem Statement"
slide_2.placeholders[1].text = (
    "🚗 No real-time parking visibility\n"
    "🔎 Drivers waste time searching\n"
    "📊 Limited parking utilization data\n"
    "⏰ Peak demand is difficult to anticipate\n"
    "🔮 Existing systems rarely provide demand prediction\n\n"
    "Our Problem:\n"
    "How can IoT and Machine Learning improve real-time parking monitoring and future demand prediction?"
)

# Slide 3: Proposed Solution & Objectives
slide_3 = prs.slides.add_slide(prs.slide_layouts[1])
slide_3.shapes.title.text = "Proposed Solution & Objectives"
slide_3.placeholders[1].text = (
    "Proposed Solution\n"
    "Distributed IoT sensors → Real-time data → Analytics → ML prediction → Dashboard\n\n"
    "Objectives\n"
    "• Monitor parking occupancy in real time\n"
    "• Transmit sensor data using IoT communication\n"
    "• Store & analyze historical occupancy data\n"
    "• Predict future parking demand\n"
    "• Visualize availability, trends & predictions"
)

# Slide 4: System Architecture (With space for PNG)
slide_4 = prs.slides.add_slide(prs.slide_layouts[1])
slide_4.shapes.title.text = "System Architecture"
slide_4.placeholders[1].text = (
    "\n\n\n[ PASTE ARCHITECTURE DIAGRAM HERE ]\n\n\n"
    "Sensing → Communication → Processing → Analytics → Prediction → Visualization"
)
for paragraph in slide_4.placeholders[1].text_frame.paragraphs:
    paragraph.alignment = 2 

# Slide 5: Data Science & ML
slide_5 = prs.slides.add_slide(prs.slide_layouts[1])
slide_5.shapes.title.text = "Data Science & Machine Learning"
slide_5.placeholders[1].text = (
    "Data\nOccupancy • Timestamp • Bay ID • Availability\n↓\n"
    "Processing\nCleaning • Aggregation • Feature Engineering\n↓\n"
    "Features\nCurrent Occupancy • Time • Day • Historical Trends\n↓\n"
    "ML Model\nRandom Forest Regression\n↓\n"
    "Prediction\nFuture Occupancy • Available Spaces • Demand"
)
for paragraph in slide_5.placeholders[1].text_frame.paragraphs:
    paragraph.alignment = 2

# Slide 6: Expected Outcome & Implementation
slide_6 = prs.slides.add_slide(prs.slide_layouts[1])
slide_6.shapes.title.text = "Expected Outcome & Implementation"
slide_6.placeholders[1].text = (
    "Expected Outcome\n"
    "• Real-Time Monitoring: Current parking availability\n"
    "• Historical Analytics: Usage & peak-demand patterns\n"
    "• ML Prediction: Future parking demand\n"
    "• Dashboard: Unified system visualization\n\n"
    "Implementation\n"
    "1. Design → 2. Build Nodes → 3. Collect Data → 4. Process → 5. Train ML → 6. Integrate Dashboard\n\n"
    "Final Goal\n"
    "Monitor parking today. Predict parking tomorrow."
)

prs.save('Smart_Parking_Presentation.pptx')
print("Presentation successfully generated: Smart_Parking_Presentation.pptx")
